from openpyxl import load_workbook
import shutil

def update_analysis_sheet(template_path, output_path, metrics):
    shutil.copyfile(template_path, output_path)

    wb = load_workbook(output_path)
    ws = wb["analysis"]

    CELL_MAP = {
        "total_rows_raw": "B1",
        "total_rows_clean": "B2",
        "rows_removed": "B3",
        "rows_removed_pct": "B4",
        "total_columns": "B5",
        "total_cells": "B6",
        "total_issues": "B7",
        "issues_per_1000_rows": "B8",
        "health_ratio": "B9",
        "auto_fix_ratio": "B10",
    }

    for key, cell in CELL_MAP.items():
        ws[cell].value = metrics.get(key, 0)

    wb.save(output_path)
